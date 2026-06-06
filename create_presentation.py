"""
SQL Server Performance Optimization - PowerPoint Generator
Creates a professional presentation with a light blue theme
Requires: pip install python-pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

def create_presentation():
    # Create presentation object
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Define light blue theme colors
    DARK_BLUE = RGBColor(41, 98, 255)      # #2962FF - Primary
    LIGHT_BLUE = RGBColor(100, 181, 246)   # #64B5F6 - Secondary
    VERY_LIGHT_BLUE = RGBColor(227, 242, 253)  # #E3F2FD - Background
    WHITE = RGBColor(255, 255, 255)
    DARK_GRAY = RGBColor(66, 66, 66)       # #424242 - Text
    
    def add_title_slide(title, subtitle):
        """Create title slide"""
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
        
        # Background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = VERY_LIGHT_BLUE
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(1))
        title_frame = title_box.text_frame
        title_frame.text = title
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(44)
        title_para.font.bold = True
        title_para.font.color.rgb = DARK_BLUE
        title_para.alignment = PP_ALIGN.CENTER
        
        # Subtitle
        subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(3.7), Inches(8), Inches(0.8))
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.text = subtitle
        subtitle_para = subtitle_frame.paragraphs[0]
        subtitle_para.font.size = Pt(24)
        subtitle_para.font.color.rgb = DARK_GRAY
        subtitle_para.alignment = PP_ALIGN.CENTER
        
        return slide
    
    def add_content_slide(title, content_lines):
        """Create content slide with bullets"""
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
        
        # Background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = WHITE
        
        # Header bar
        header = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1))
        header.fill.solid()
        header.fill.fore_color.rgb = LIGHT_BLUE
        header.line.fill.background()
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.6))
        title_frame = title_box.text_frame
        title_frame.text = title
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(32)
        title_para.font.bold = True
        title_para.font.color.rgb = WHITE
        
        # Content
        content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(8.6), Inches(5.7))
        text_frame = content_box.text_frame
        text_frame.word_wrap = True
        
        for i, line in enumerate(content_lines):
            if i > 0:
                p = text_frame.add_paragraph()
            else:
                p = text_frame.paragraphs[0]
            
            # Determine indentation level
            indent_level = 0
            clean_line = line.lstrip()
            if line.startswith('  - ') or line.startswith('    •'):
                indent_level = 1
                clean_line = clean_line[2:].strip()
            elif line.startswith('- ') or line.startswith('• '):
                indent_level = 0
                clean_line = clean_line[2:].strip()
            
            p.text = clean_line
            p.level = indent_level
            p.font.size = Pt(16) if indent_level == 0 else Pt(14)
            p.font.color.rgb = DARK_GRAY
            p.space_before = Pt(6)
            
        return slide
    
    def add_code_slide(title, code_text, description=None):
        """Create slide with code example"""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        
        # Background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = WHITE
        
        # Header
        header = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1))
        header.fill.solid()
        header.fill.fore_color.rgb = LIGHT_BLUE
        header.line.fill.background()
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.6))
        title_frame = title_box.text_frame
        title_frame.text = title
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(32)
        title_para.font.bold = True
        title_para.font.color.rgb = WHITE
        
        # Description if provided
        y_pos = 1.3
        if description:
            desc_box = slide.shapes.add_textbox(Inches(0.7), Inches(y_pos), Inches(8.6), Inches(0.6))
            desc_frame = desc_box.text_frame
            desc_frame.text = description
            desc_para = desc_frame.paragraphs[0]
            desc_para.font.size = Pt(14)
            desc_para.font.color.rgb = DARK_GRAY
            y_pos += 0.8
        
        # Code box
        code_box = slide.shapes.add_textbox(Inches(0.7), Inches(y_pos), Inches(8.6), Inches(6.5 - y_pos))
        code_frame = code_box.text_frame
        code_frame.text = code_text
        code_frame.word_wrap = True
        
        for paragraph in code_frame.paragraphs:
            paragraph.font.name = 'Consolas'
            paragraph.font.size = Pt(12)
            paragraph.font.color.rgb = DARK_GRAY
        
        # Code background
        code_box.fill.solid()
        code_box.fill.fore_color.rgb = RGBColor(245, 245, 245)
        
        return slide
    
    # SLIDE 1: Title Slide
    add_title_slide(
        "SQL Server Performance Optimization",
        "Indexes, Execution Plans, and Query Optimization"
    )
    
    # SLIDE 2: Agenda
    add_content_slide("Today's Topics", [
        "1. Index Structures and Maintenance",
        "2. Reading and Optimizing Execution Plans",
        "3. Advanced Index Techniques",
        "4. Query Processing Order",
        "5. Implicit Conversions",
        "6. Query Anti-Patterns to Avoid",
        "7. Query Best Practices",
        "8. Table-Valued Functions",
        "9. Parameter Sniffing Solutions"
    ])
    
    # SLIDE 3: Index Structures
    add_content_slide("Index Storage Structures", [
        "Heap Tables: No clustered index",
        "  - Fast inserts, slow reads",
        "  - Best for staging tables, small lookup tables",
        "B+-Tree: SQL Server's index implementation",
        "  - All data stored in leaf nodes",
        "  - 3-4 levels for millions of rows",
        "  - Efficient range scans via linked leaf nodes",
        "  - Better cache utilization"
    ])
    
    # SLIDE 4: B+-Tree Diagram
    add_content_slide("How B+-Tree Indexes Work", [
        "Structure:",
        "  - Root node at top",
        "  - Intermediate nodes (keys and pointers only)",
        "  - Leaf nodes (all data, doubly-linked)",
        "",
        "Key advantages:",
        "  - Balanced structure (all paths same length)",
        "  - Sequential access at leaf level",
        "  - Typical depth: 3-4 levels for millions of rows"
    ])
    
    # SLIDE 5: Fragmentation
    add_content_slide("Understanding Index Fragmentation", [
        "Logical Fragmentation:",
        "  - Pages not physically sequential",
        "  - Measured: avg_fragmentation_in_percent",
        "  - Causes: INSERTs, UPDATEs, DELETEs",
        "  - Impact: More physical I/O",
        "",
        "Page Density: avg_page_space_used_in_percent",
        "  - Low density = wasted space",
        "  - High density = good utilization"
    ])
    
    # SLIDE 6: Rebuild vs Reorganize
    add_content_slide("Index Maintenance Strategies", [
        "Decision Matrix:",
        "  - < 10% fragmentation: Do Nothing",
        "  - 10-30% fragmentation: REORGANIZE",
        "  - > 30% fragmentation: REBUILD",
        "",
        "REORGANIZE: Online, minimal locks, compacts pages",
        "REBUILD: Complete defrag, updates statistics, needs space"
    ])
    
    # SLIDE 7: Demo 1
    add_content_slide("DEMO: Index Fragmentation", [
        "What we'll show:",
        "  - Check fragmentation levels",
        "  - Create fragmentation through random inserts",
        "  - Compare Rebuild vs. Reorganize",
        "  - Show performance improvement",
        "  - DMV: sys.dm_db_index_physical_stats"
    ])
    
    # SLIDE 8: Reading Execution Plans
    add_content_slide("Reading Execution Plans", [
        "Direction: Right to left, top to bottom",
        "Thick arrows: More rows flowing through",
        "Color coding:",
        "  - Black/Blue: Normal operation",
        "  - Yellow triangle: Warnings",
        "Relative cost: Resource consumption percentages",
        "Focus on: Highest cost operators"
    ])
    
    # SLIDE 9: Key Operators
    add_content_slide("Common Execution Plan Operators", [
        "Scans (Read everything):",
        "  - Table Scan, Index Scan, Clustered Index Scan",
        "Seeks (Targeted read - GOOD):",
        "  - Index Seek, Clustered Index Seek",
        "Joins:",
        "  - Nested Loop (small datasets)",
        "  - Hash Match (large unsorted data)",
        "  - Merge Join (large sorted data)",
        "Key Lookup: Additional data retrieval (often costly)"
    ])
    
    # SLIDE 10: Plan Warnings
    add_content_slide("Interpreting Plan Warnings", [
        "Yellow exclamation marks indicate:",
        "  - Implicit type conversions",
        "  - Missing statistics",
        "  - Excessive memory grants",
        "  - Spills to tempdb",
        "",
        "Missing Index hints (green text):",
        "  - Take with caution",
        "  - Evaluate cost vs. benefit",
        "  - May suggest too many indexes"
    ])
    
    # SLIDE 11: Optimization Strategies
    add_content_slide("How to Optimize Based on Plans", [
        "1. Eliminate Scans → Add appropriate indexes",
        "2. Reduce Key Lookups → Covering indexes",
        "3. Avoid Sorts → Index matches ORDER BY",
        "4. Fix Conversions → Match data types",
        "5. Update Statistics → Accurate cardinality",
        "6. Rewrite Logic → Sometimes better than indexes"
    ])
    
    # SLIDE 12: Demo 2
    add_content_slide("DEMO: Execution Plans", [
        "What we'll demonstrate:",
        "  - Compare Scan vs. Seek performance",
        "  - Show Key Lookup overhead",
        "  - Different JOIN operators in action",
        "  - Fix a slow query using execution plan",
        "  - Before/After metrics (IO, time)"
    ])
    
    # SLIDE 13: Index Types
    add_content_slide("Clustered vs. Nonclustered Indexes", [
        "Clustered Index:",
        "  - Determines physical row order",
        "  - One per table, leaf level = actual data",
        "  - Choose: Narrow, unique, static, increasing",
        "",
        "Nonclustered Index:",
        "  - Separate structure from data",
        "  - Up to 999 per table (don't!)",
        "  - Leaf level = key + bookmark",
        "  - Overhead on INSERT/UPDATE/DELETE"
    ])
    
    # SLIDE 14: Covering Indexes
    add_code_slide(
        "Covering Indexes with INCLUDE",
        "CREATE NONCLUSTERED INDEX IX_Example\nON Orders (CustomerID, OrderDate)\nINCLUDE (Amount, Status);",
        "Benefits: Query satisfied from index, no bookmark lookup needed"
    )
    
    # SLIDE 15: Filtered Indexes
    add_code_slide(
        "Indexes on Subsets of Data",
        "CREATE NONCLUSTERED INDEX IX_ActiveOrders\nON Orders (OrderDate, CustomerID)\nWHERE Status = 'Active';",
        "Use for: Sparse columns, status-based filtering, current year data"
    )
    
    # SLIDE 16: Demo 3
    add_content_slide("DEMO: Index Types", [
        "What we'll show:",
        "  - Create covering index",
        "  - Eliminate Key Lookups",
        "  - Filtered index benefits",
        "  - Compare index sizes",
        "  - Index usage statistics (DMVs)"
    ])
    
    # SLIDE 17: Query Processing Order
    add_content_slide("Logical Query Processing Order", [
        "1. FROM (including JOINs)",
        "2. WHERE",
        "3. GROUP BY",
        "4. HAVING",
        "5. SELECT",
        "6. DISTINCT",
        "7. ORDER BY",
        "8. TOP / OFFSET-FETCH",
        "",
        "Key Point: Logical ≠ Physical execution order"
    ])
    
    # SLIDE 18: Order Mistakes
    add_code_slide(
        "Common Mistakes - Processing Order",
        "-- WRONG\nSELECT Price * Quantity AS Total\nFROM Orders\nWHERE Total > 1000;  -- Error!\n\n-- CORRECT\nSELECT Price * Quantity AS Total\nFROM Orders\nWHERE Price * Quantity > 1000;",
        "Why? WHERE executes before SELECT"
    )
    
    # SLIDE 19: WHERE vs HAVING
    add_code_slide(
        "Filter Early with WHERE",
        "SELECT CustomerID, COUNT(*) AS OrderCount\nFROM Orders\nWHERE OrderDate >= '2024-01-01'  -- Filter early\nGROUP BY CustomerID\nHAVING COUNT(*) > 5;  -- Filter groups",
        "WHERE: Before aggregation (efficient) | HAVING: After aggregation"
    )
    
    # SLIDE 20: Implicit Conversions
    add_code_slide(
        "The Hidden Performance Killer",
        "-- BAD: NVARCHAR literal on VARCHAR column\nWHERE Username = N'john'  -- Index scan\n\n-- GOOD: Matching types\nWHERE Username = 'john'   -- Index seek",
        "Problem: Auto-conversion prevents index usage"
    )
    
    # SLIDE 21: Conversion Scenarios
    add_content_slide("Where Conversions Hide", [
        "1. VARCHAR vs NVARCHAR: N prefix on literals",
        "2. Numeric precision: INT vs DECIMAL/FLOAT",
        "3. JOIN mismatches: INT vs BIGINT",
        "4. Date functions: CONVERT(DATE, column)",
        "",
        "Impact:",
        "  - Index seeks become scans",
        "  - CPU overhead, inaccurate cardinality",
        "  - Potential query timeouts"
    ])
    
    # SLIDE 22: Demo 4
    add_content_slide("DEMO: Implicit Conversions", [
        "What we'll demonstrate:",
        "  - Show conversion warning in plan",
        "  - Performance comparison: with vs without",
        "  - Fix JOIN type mismatch",
        "  - DMV query to find conversions",
        "  - Before/After logical reads"
    ])
    
    # SLIDE 23: Anti-Patterns
    add_content_slide("Worst Patterns That Kill Performance", [
        "1. Functions on columns in WHERE",
        "2. Leading wildcards in LIKE",
        "3. OR across different columns",
        "4. RBAR (Row-By-Agonizing-Row)",
        "5. SELECT *",
        "6. NOT IN with NULLs",
        "7. Correlated subqueries in SELECT"
    ])
    
    # SLIDE 24: Functions in WHERE
    add_code_slide(
        "Anti-Pattern: Functions in WHERE",
        "-- BAD\nWHERE YEAR(OrderDate) = 2024\nWHERE UPPER(LastName) = 'SMITH'\n\n-- GOOD (Sargable)\nWHERE OrderDate >= '2024-01-01' \n  AND OrderDate < '2025-01-01'\nWHERE LastName = 'Smith'",
        "Result: Index Seek instead of Scan"
    )
    
    # SLIDE 25: RBAR
    add_code_slide(
        "Anti-Pattern: Row-By-Agonizing-Row",
        "-- BAD: Cursors\nDECLARE cur CURSOR FOR SELECT CustomerID...\nWHILE @@FETCH_STATUS = 0\nBEGIN\n    UPDATE Orders SET Status = 'Processed' ...\nEND\n\n-- GOOD: Set-based\nUPDATE Orders SET Status = 'Processed'\nFROM Orders O\nJOIN Customers C ON O.CustomerID = C.CustomerID;",
        "Set-based operations are orders of magnitude faster"
    )
    
    # SLIDE 26: Demo 5
    add_content_slide("DEMO: Anti-Patterns", [
        "What we'll show:",
        "  - Function in WHERE: Before/After",
        "  - RBAR vs Set-based comparison",
        "  - SELECT * vs specific columns",
        "  - NOT IN vs NOT EXISTS with NULLs",
        "  - Execution time differences"
    ])
    
    # SLIDE 27: Best Practices
    add_content_slide("Query Best Patterns", [
        "1. Filter early and often",
        "2. EXISTS vs IN (EXISTS stops at first match)",
        "3. CTEs for readability",
        "4. Window functions vs self-joins",
        "5. CROSS APPLY for top N per group",
        "6. Batch updates for large data",
        "7. Covering indexes with INCLUDE"
    ])
    
    # SLIDE 28: CTEs
    add_code_slide(
        "Best Practice: Common Table Expressions",
        "WITH ActiveCustomers AS (\n    SELECT CustomerID, CustomerName\n    FROM Customers WHERE Status = 'Active'\n),\nRecentOrders AS (\n    SELECT CustomerID, OrderDate, Amount\n    FROM Orders\n    WHERE OrderDate >= DATEADD(MONTH, -3, GETDATE())\n)\nSELECT AC.CustomerName, COUNT(*) AS OrderCount\nFROM ActiveCustomers AC\nJOIN RecentOrders RO ON AC.CustomerID = RO.CustomerID\nGROUP BY AC.CustomerName;",
        "Benefits: Readable, maintainable, easier debugging"
    )
    
    # SLIDE 29: Window Functions
    add_code_slide(
        "Best Practice: Window Functions",
        "-- OLD (Slow)\nSELECT OrderDate, \n  (SELECT SUM(Amount) FROM Orders O2 \n   WHERE O2.OrderDate <= O1.OrderDate) AS RunningTotal\nFROM Orders O1;\n\n-- BETTER (Fast)\nSELECT OrderDate,\n  SUM(Amount) OVER (ORDER BY OrderDate \n    ROWS UNBOUNDED PRECEDING) AS RunningTotal\nFROM Orders;",
        "Single pass vs multiple self-joins"
    )
    
    # SLIDE 30: Demo 6
    add_content_slide("DEMO: Best Practices", [
        "What we'll demonstrate:",
        "  - CTE readability example",
        "  - Window function performance",
        "  - CROSS APPLY for top N per group",
        "  - Batch update vs single update",
        "  - Covering index benefits"
    ])
    
    # SLIDE 31: TVF Comparison
    add_content_slide("Table-Valued Functions Comparison", [
        "Inline TVF:",
        "  - Excellent performance",
        "  - Fully optimized by query processor",
        "  - Accurate cardinality, supports parallelism",
        "  - Use: Almost always!",
        "",
        "Multi-Statement TVF:",
        "  - Poor performance (black box)",
        "  - Fixed cardinality estimate (100 rows)",
        "  - No parallelism, no statistics",
        "  - Use: Only if procedural logic required"
    ])
    
    # SLIDE 32: Inline TVF
    add_code_slide(
        "The Right Way: Inline TVFs",
        "CREATE FUNCTION dbo.GetCustomerOrders(@CustomerID INT)\nRETURNS TABLE\nAS\nRETURN\n(\n    SELECT OrderID, OrderDate, Amount\n    FROM Orders\n    WHERE CustomerID = @CustomerID\n);\n\n-- Usage\nSELECT * FROM dbo.GetCustomerOrders(123);",
        "Acts like a parameterized view - optimizer can see through it!"
    )
    
    # SLIDE 33: Demo 7
    add_content_slide("DEMO: Table-Valued Functions", [
        "What we'll show:",
        "  - Create inline TVF",
        "  - Create multi-statement TVF",
        "  - Performance comparison (same logic)",
        "  - Execution plan differences",
        "  - Scalar UDF anti-pattern (avoid!)"
    ])
    
    # SLIDE 34: Parameter Sniffing
    add_content_slide("What is Parameter Sniffing?", [
        "SQL Server behavior:",
        "  - Compiles proc with first parameter values",
        "  - Plan cached and reused for all executions",
        "",
        "The Good: Efficient plan reuse",
        "The Bad: Plan optimal for first, terrible for others",
        "",
        "Example:",
        "  - First: @Status = 'Pending' (1000 rows) → Scan",
        "  - Second: @Status = 'Cancelled' (3 rows) → Uses Scan!"
    ])
    
    # SLIDE 35: Detection
    add_content_slide("Detecting Parameter Sniffing", [
        "Symptoms:",
        "  - Same procedure fast sometimes, slow others",
        "  - Performance varies with parameter values",
        "  - Same plan for different row counts",
        "",
        "Detection methods:",
        "  - Execution plan: same shape for different params",
        "  - DMV: sys.dm_exec_query_stats (variance)",
        "  - Extended Events: parameter_sensitive_plan"
    ])
    
    # SLIDE 36: Solutions 1-2
    add_code_slide(
        "Parameter Sniffing Solutions (1/2)",
        "-- 1. OPTIMIZE FOR hint\nOPTION (OPTIMIZE FOR (@Status = 'Pending'))\n-- Use when: One value is 90%+ of calls\n\n-- 2. OPTIMIZE FOR UNKNOWN\nOPTION (OPTIMIZE FOR UNKNOWN)\n-- Use when: Even distribution of values",
        "Choose based on your data distribution"
    )
    
    # SLIDE 37: Solutions 3-7
    add_content_slide("Parameter Sniffing Solutions (2/2)", [
        "3. RECOMPILE",
        "  - WITH RECOMPILE or OPTION (RECOMPILE)",
        "  - Use: High sensitivity, low frequency",
        "",
        "4. Local variable copy",
        "  - DECLARE @Local = @Parameter",
        "",
        "5. Dynamic SQL (fresh plan each time)",
        "6. Multiple procedures (branch logic)",
        "7. SQL 2022: Parameter Sensitive Plans (PSP) - Automatic!"
    ])
    
    # SLIDE 38: Decision Matrix
    add_content_slide("Choosing the Right Solution", [
        "Decision Matrix:",
        "  - One value dominates (90%+) → OPTIMIZE FOR value",
        "  - Even distribution → OPTIMIZE FOR UNKNOWN",
        "  - Extreme sensitivity → RECOMPILE",
        "  - Ad-hoc reporting → Dynamic SQL",
        "  - SQL Server 2022+ → Enable PSP",
        "",
        "SQL 2022 PSP: Automatically detects and creates multiple plans!"
    ])
    
    # SLIDE 39: Demo 8
    add_content_slide("DEMO: Parameter Sniffing", [
        "What we'll demonstrate:",
        "  - Show problem: Same plan, different data",
        "  - OPTIMIZE FOR solution",
        "  - OPTIMIZE FOR UNKNOWN comparison",
        "  - RECOMPILE overhead measurement",
        "  - Performance metrics comparison"
    ])
    
    # SLIDE 40: Metrics
    add_content_slide("Key Performance Metrics", [
        "Always monitor:",
        "  - Logical reads (SET STATISTICS IO)",
        "  - CPU time vs Elapsed time",
        "  - Execution plan warnings",
        "  - Index fragmentation percentage",
        "  - Plan cache reuse",
        "",
        "Essential DMVs:",
        "  - sys.dm_db_index_physical_stats",
        "  - sys.dm_exec_query_stats",
        "  - sys.dm_db_index_usage_stats"
    ])
    
    # SLIDE 41: Interview Questions
    add_content_slide("Common Interview Questions", [
        "1. When use clustered vs nonclustered index?",
        "2. How identify a missing index?",
        "3. Difference between REBUILD and REORGANIZE?",
        "4. How implicit conversions affect performance?",
        "5. When use a filtered index?",
        "6. What is parameter sniffing and solutions?",
        "7. What makes a query 'sargable'?"
    ])
    
    # SLIDE 42: Best Practices Summary
    add_content_slide("Key Takeaways", [
        "✓ Monitor fragmentation - Rebuild > 30%, Reorganize 10-30%",
        "✓ Read execution plans - Focus on scans, lookups, warnings",
        "✓ Use covering indexes - INCLUDE for non-key columns",
        "✓ Filter early - WHERE before joins",
        "✓ Avoid functions on columns - Use sargable predicates",
        "✓ Prefer inline TVFs - Never scalar UDFs",
        "✓ Match data types - Prevent implicit conversions",
        "✓ Handle parameter sniffing - Choose appropriate solution"
    ])
    
    # SLIDE 43: Tools and Resources
    add_content_slide("Tools and Resources", [
        "Tools:",
        "  - SQL Server Management Studio (SSMS)",
        "  - Extended Events",
        "  - Database Tuning Advisor",
        "  - Query Store (SQL 2016+)",
        "",
        "Documentation:",
        "  - Microsoft SQL Server Docs",
        "  - Execution Plan Reference",
        "  - Index Design Guide"
    ])
    
    # SLIDE 44: Demo Code
    add_content_slide("Hands-On Practice", [
        "Demo Code Repository:",
        "  - All demo scripts available",
        "  - 9 comprehensive SQL files",
        "  - Sample database creation",
        "  - Performance comparison examples",
        "  - README with setup instructions",
        "",
        "Database: PerformanceDemo",
        "Files: 01-09 covering all topics"
    ])
    
    # SLIDE 45: Q&A
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = LIGHT_BLUE
    
    title_box = slide.shapes.add_textbox(Inches(2), Inches(2.5), Inches(6), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.text = "Questions?"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(60)
    title_para.font.bold = True
    title_para.font.color.rgb = WHITE
    title_para.alignment = PP_ALIGN.CENTER
    
    subtitle_box = slide.shapes.add_textbox(Inches(2), Inches(4.2), Inches(6), Inches(0.8))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "Open Discussion • Real-World Scenarios"
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(20)
    subtitle_para.font.color.rgb = WHITE
    subtitle_para.alignment = PP_ALIGN.CENTER
    
    # SLIDE 46: Thank You
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BLUE
    
    title_box = slide.shapes.add_textbox(Inches(1.5), Inches(2), Inches(7), Inches(1))
    title_frame = title_box.text_frame
    title_frame.text = "Thank You!"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(54)
    title_para.font.bold = True
    title_para.font.color.rgb = WHITE
    title_para.alignment = PP_ALIGN.CENTER
    
    message_box = slide.shapes.add_textbox(Inches(1.5), Inches(3.5), Inches(7), Inches(2))
    message_frame = message_box.text_frame
    message_frame.text = "Performance optimization is an ongoing process:\nMonitor • Test • Update Statistics • Review Plans\n\nThe best optimization solves YOUR specific problem!"
    message_frame.word_wrap = True
    for paragraph in message_frame.paragraphs:
        paragraph.font.size = Pt(18)
        paragraph.font.color.rgb = WHITE
        paragraph.alignment = PP_ALIGN.CENTER
    
    # Save presentation
    output_file = 'SQL_Server_Performance_Optimization.pptx'
    prs.save(output_file)
    print(f"✓ Presentation created: {output_file}")
    print(f"✓ Total slides: {len(prs.slides)}")
    print(f"✓ Theme: Light Blue")
    
    return output_file

if __name__ == "__main__":
    print("Creating SQL Server Performance Optimization presentation...")
    print("-" * 60)
    
    try:
        filename = create_presentation()
        print("-" * 60)
        print(f"Success! Open '{filename}' in PowerPoint.")
    except ImportError:
        print("\nERROR: python-pptx library not found!")
        print("Please install it using:")
        print("  pip install python-pptx")
    except Exception as e:
        print(f"\nERROR: {e}")
